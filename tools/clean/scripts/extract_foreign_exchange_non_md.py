from __future__ import annotations

import argparse
import hashlib
import json
import tempfile
from collections import Counter
from pathlib import Path
from typing import Callable, Iterable


NON_MARKDOWN_EXTENSIONS = {".doc", ".docx", ".pdf", ".xls", ".xlsx"}


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def write_jsonl(path: Path, rows: Iterable[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as fh:
        for row in rows:
            fh.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def clean_output_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)
    for child in path.iterdir():
        if child.is_file():
            child.unlink()


def markdown_path_for(parsed_dir: Path, row: dict, used_names: set[str]) -> Path:
    stem = Path(row["file_name"]).stem
    ext_suffix = row.get("extension", "").lstrip(".")
    candidate = f"{stem}.md"
    if candidate in used_names:
        candidate = f"{stem}__{ext_suffix}.md"
    counter = 2
    while candidate in used_names:
        candidate = f"{stem}__{ext_suffix}_{counter}.md"
        counter += 1
    used_names.add(candidate)
    return parsed_dir / candidate


def compact_text(text: str) -> str:
    return "\n".join(line.rstrip() for line in text.splitlines()).strip()


def count_non_empty_lines(text: str) -> int:
    return sum(1 for line in text.splitlines() if line.strip())


def wrap_markdown(title: str, body: str) -> str:
    body = compact_text(body)
    if not body:
        return ""
    return f"# {Path(title).stem}\n\n{body}\n"


def extract_docx_text(path: Path) -> tuple[str, dict]:
    from docx import Document

    document = Document(path)
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    table_count = 0
    for table in document.tables:
        table_count += 1
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n\n".join(parts), {"table_count": table_count, "page_count": 0, "worksheet_count": 0}


def convert_doc_with_word(path: Path, temp_dir: Path) -> Path:
    import win32com.client  # type: ignore

    output = temp_dir / f"{path.stem}.docx"
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    document = None
    try:
        document = word.Documents.Open(str(path), ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False)
        document.SaveAs2(str(output), FileFormat=16)
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()
    return output


def extract_doc_text(path: Path, converter: Callable[[Path, Path], Path] | None = None) -> tuple[str, dict]:
    converter = converter or convert_doc_with_word
    with tempfile.TemporaryDirectory() as td:
        converted = converter(path, Path(td))
        text, metrics = extract_docx_text(converted)
        metrics["converted_from_doc"] = True
        return text, metrics


def extract_pdf_text(path: Path) -> tuple[str, dict]:
    import pdfplumber

    parts: list[str] = []
    page_count = 0
    table_count = 0
    with pdfplumber.open(path) as pdf:
        page_count = len(pdf.pages)
        for idx, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"## Page {idx}\n\n{text.strip()}")
            for table in page.extract_tables() or []:
                table_count += 1
                for row in table:
                    cells = [str(cell).strip().replace("\n", " ") if cell is not None else "" for cell in row]
                    if any(cells):
                        parts.append(" | ".join(cells))
    return "\n\n".join(parts), {"table_count": table_count, "page_count": page_count, "worksheet_count": 0}


def extract_xlsx_text(path: Path) -> tuple[str, dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []
    worksheet_count = 0
    row_count = 0
    for sheet in workbook.worksheets:
        worksheet_count += 1
        parts.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value).strip() for value in row]
            if any(values):
                row_count += 1
                parts.append(" | ".join(values))
    workbook.close()
    return "\n\n".join(parts), {"table_count": row_count, "page_count": 0, "worksheet_count": worksheet_count}


def extract_xls_text(path: Path) -> tuple[str, dict]:
    import xlrd

    workbook = xlrd.open_workbook(path)
    parts: list[str] = []
    row_count = 0
    for sheet in workbook.sheets():
        parts.append(f"## Sheet: {sheet.name}")
        for row_idx in range(sheet.nrows):
            values = [str(sheet.cell_value(row_idx, col_idx)).strip() for col_idx in range(sheet.ncols)]
            if any(values):
                row_count += 1
                parts.append(" | ".join(values))
    return "\n\n".join(parts), {"table_count": row_count, "page_count": 0, "worksheet_count": workbook.nsheets}


def empty_stats() -> dict:
    return {"table_count": 0, "page_count": 0, "worksheet_count": 0}


def build_failure_record(row: dict, status: str, reason: str, parser: str) -> dict:
    return {
        "source_path": row["source_path"],
        "file_name": row["file_name"],
        "extension": row["extension"],
        "size_bytes": row.get("size_bytes", 0),
        "is_form_like_suspected": bool(row.get("is_form_like_suspected", False)),
        "is_duplicate_version_suspected": bool(row.get("is_duplicate_version_suspected", False)),
        "parse_status": status,
        "review_reason": reason,
        "parser": parser,
        "parsed_path": None,
        "text_sha256": None,
        "char_count": 0,
        "non_empty_line_count": 0,
        **empty_stats(),
    }


def build_success_record(row: dict, parser: str, parsed_path: Path, text: str, metrics: dict) -> dict:
    return {
        "source_path": row["source_path"],
        "file_name": row["file_name"],
        "extension": row["extension"],
        "size_bytes": row.get("size_bytes", 0),
        "is_form_like_suspected": bool(row.get("is_form_like_suspected", False)),
        "is_duplicate_version_suspected": bool(row.get("is_duplicate_version_suspected", False)),
        "parse_status": "parsed",
        "review_reason": None,
        "parser": parser,
        "parsed_path": str(parsed_path),
        "text_sha256": sha256_text(text),
        "char_count": len(text),
        "non_empty_line_count": count_non_empty_lines(text),
        "table_count": int(metrics.get("table_count", 0)),
        "page_count": int(metrics.get("page_count", 0)),
        "worksheet_count": int(metrics.get("worksheet_count", 0)),
    }


def extract_one(
    row: dict,
    parsed_dir: Path,
    used_output_names: set[str],
    doc_converter: Callable[[Path, Path], Path] | None = None,
) -> dict:
    path = Path(row["source_path"])
    ext = row["extension"].lower()
    parser = ext.lstrip(".")
    try:
        if ext == ".docx":
            body, metrics = extract_docx_text(path)
        elif ext == ".doc":
            body, metrics = extract_doc_text(path, converter=doc_converter)
            parser = "word_com_doc"
        elif ext == ".pdf":
            body, metrics = extract_pdf_text(path)
        elif ext == ".xlsx":
            body, metrics = extract_xlsx_text(path)
        elif ext == ".xls":
            body, metrics = extract_xls_text(path)
        else:
            return build_failure_record(row, "failed", f"unsupported_extension:{ext}", parser)
    except Exception as exc:
        reason = "doc_conversion_failed" if ext == ".doc" else "parse_failed"
        return build_failure_record(row, "needs_review", f"{reason}: {exc}", parser)

    text = wrap_markdown(row["file_name"], body)
    if not text.strip():
        reason = "pdf_no_extractable_text" if ext == ".pdf" else "empty_extracted_text"
        return build_failure_record(row, "needs_review", reason, parser)

    parsed_path = markdown_path_for(parsed_dir, row, used_output_names)
    parsed_path.write_text(text, encoding="utf-8", newline="\n")
    return build_success_record(row, parser, parsed_path, text, metrics)


def summarize(records: list[dict], manifest_path: Path, output_dir: Path) -> dict:
    status_counts = Counter(row["parse_status"] for row in records)
    ext_counts = Counter(row["extension"] for row in records)
    return {
        "manifest": str(manifest_path),
        "parse_report": str(output_dir / "parse_report.jsonl"),
        "candidate_count": len(records),
        "status_counts": dict(sorted(status_counts.items())),
        "extension_counts": dict(sorted(ext_counts.items())),
        "parsed_count": status_counts.get("parsed", 0),
        "needs_review_count": status_counts.get("needs_review", 0),
        "failed_count": status_counts.get("failed", 0),
    }


def extract_non_markdown_corpus(
    manifest_path: Path | str,
    output_dir: Path | str,
    doc_converter: Callable[[Path, Path], Path] | None = None,
) -> dict:
    manifest = Path(manifest_path)
    output = Path(output_dir)
    parsed_dir = output / "parsed_text"
    output.mkdir(parents=True, exist_ok=True)
    clean_output_dir(parsed_dir)

    rows = [
        row
        for row in read_jsonl(manifest)
        if not row.get("is_markdown", False) and row.get("extension", "").lower() in NON_MARKDOWN_EXTENSIONS
    ]
    used_output_names: set[str] = set()
    records = [
        extract_one(row, parsed_dir, used_output_names, doc_converter=doc_converter)
        for row in rows
    ]
    write_jsonl(output / "parse_report.jsonl", records)
    summary = summarize(records, manifest, output)
    (output / "parse_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
        newline="\n",
    )
    return summary


def main() -> None:
    clean_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Extract text from non-Markdown foreign exchange files.")
    parser.add_argument("--manifest", default=str(clean_root / "manifest" / "manifest.jsonl"))
    parser.add_argument("--output", default=str(clean_root / "non_md_parse"))
    args = parser.parse_args()

    summary = extract_non_markdown_corpus(args.manifest, args.output)
    print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
