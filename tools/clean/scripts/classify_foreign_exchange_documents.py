from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import tempfile
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable


SOURCE_DIR = Path(r"E:\外汇文件_filter_分组")
DEFAULT_CSV = Path("docs/eval/foreign_exchange_document_classification.csv")
DEFAULT_SUMMARY = Path("docs/eval/foreign_exchange_document_classification_summary.md")


FORM_NAME_TERMS = (
    "申请表",
    "申请单",
    "信息表",
    "变更表",
    "退市申请表",
    "材料清单",
    "要素表",
    "名单",
    "基本信息表",
    "登记表",
    "备案申请表",
    "委托书",
    "application form",
)

GENERAL_NAME_TERMS = (
    "服务介绍",
    "服务和标准",
    "产品介绍",
    "日历",
    "清单",
    "要素表",
    "合约要素",
    "信息表",
    "申请表",
    "申请单",
    "变更表",
    "名单",
    "材料清单",
    "说明",
    "基本信息表",
    "委托书",
    "application form",
)

MANUAL_STRONG_NAME_TERMS = (
    "操作手册",
    "操作指南",
    "使用手册",
    "入市指南",
    "开户指南",
    "调整指南",
    "快速使用指南",
    "流程指南",
    "联网开户",
    "开户流程",
    "配置要求",
    "投资者问答",
    "问答",
    "market entry guide",
    "operational guide",
    "account opening process",
    "networking and account opening",
)

MANUAL_WEAK_NAME_TERMS = (
    "手册",
    "指南",
    "流程",
    "配置",
    "使用",
    "开户",
    "联网",
    "入市",
    "guide",
)

LAWS_STRONG_NAME_TERMS = (
    "交易规则",
    "交易指引",
    "业务指引",
    "产品指引",
    "操作指引",
    "业务规范",
    "技术规范",
    "数据规范",
    "服务协议",
    "标准条款",
    "操作规程",
    "操作细则",
    "实施细则",
    "管理办法",
    "交易规程",
    "确认规则",
    "准入指引",
)

LAWS_NAME_TERMS = (
    "规则",
    "规程",
    "细则",
    "办法",
    "规范",
    "协议",
    "条款",
    "制度",
    "准入",
    "管理",
    "承诺函",
)

LAW_STRUCTURE_TERMS = (
    "第一条",
    "第二条",
    "第三条",
    "总则",
    "附则",
    "第一章",
    "第二章",
    "第三章",
    "本规则",
    "本办法",
    "本规范",
    "本协议",
    "本细则",
    "本规程",
)

LAW_NORMATIVE_TERMS = (
    "应当",
    "不得",
    "适用于",
    "负责",
    "有效期",
    "申请机构应",
    "交易成员应",
    "中心负责",
    "另有规定",
)

MANUAL_ACTION_TERMS = (
    "步骤",
    "登录",
    "点击",
    "选择",
    "填写",
    "提交",
    "下载",
    "安装",
    "配置",
    "菜单",
    "界面",
    "客户端",
    "系统",
    "用户",
    "操作",
    "开户",
    "联网",
)

NEWS_HINT_TERMS = ("发布", "上线", "扩容", "优化措施", "正式发布", "资料")


@dataclass
class ExtractResult:
    text: str
    status: str
    parser: str
    error: str = ""
    page_count: int = 0


@dataclass
class ClassifyResult:
    category: str
    confidence: str
    evidence: str
    needs_review: str
    review_reason: str


def discover_files(source_dir: Path) -> list[Path]:
    return sorted(
        [path for path in source_dir.rglob("*") if path.is_file()],
        key=lambda p: str(p).lower(),
    )


def contains_any(text: str, terms: Iterable[str]) -> bool:
    lowered = text.lower()
    return any(term.lower() in lowered for term in terms)


def matched_terms(text: str, terms: Iterable[str]) -> list[str]:
    lowered = text.lower()
    return [term for term in terms if term.lower() in lowered]


def count_terms(text: str, terms: Iterable[str]) -> int:
    lowered = text.lower()
    total = 0
    for term in terms:
        total += lowered.count(term.lower())
    return total


def compact(text: str, limit: int = 80000) -> str:
    text = re.sub(r"[ \t]+", " ", text or "")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:limit]


def read_text_with_fallback(path: Path) -> ExtractResult:
    encodings = ("utf-8-sig", "utf-8", "gb18030")
    errors: list[str] = []
    for enc in encodings:
        try:
            return ExtractResult(
                text=compact(path.read_text(encoding=enc)),
                status="parsed",
                parser=f"text:{enc}",
            )
        except Exception as exc:
            errors.append(f"{enc}: {exc}")
    return ExtractResult("", "parse_failed", "text", "; ".join(errors))


def extract_docx_text(path: Path) -> ExtractResult:
    try:
        from docx import Document

        document = Document(path)
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = compact("\n".join(parts))
        return ExtractResult(text, "parsed" if text else "empty_text", "python-docx")
    except Exception as exc:
        return ExtractResult("", "parse_failed", "python-docx", str(exc))


def convert_doc_with_word(path: Path, temp_dir: Path) -> Path:
    import win32com.client  # type: ignore

    output = temp_dir / f"{path.stem}.docx"
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    document = None
    try:
        document = word.Documents.Open(
            str(path),
            ConfirmConversions=False,
            ReadOnly=True,
            AddToRecentFiles=False,
        )
        document.SaveAs2(str(output), FileFormat=16)
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()
    return output


def extract_doc_text(path: Path) -> ExtractResult:
    try:
        with tempfile.TemporaryDirectory(prefix="openrag_doc_classify_") as td:
            converted = convert_doc_with_word(path, Path(td))
            result = extract_docx_text(converted)
            result.parser = "word-com->python-docx"
            if result.status == "parsed":
                return result
            return ExtractResult(result.text, result.status, result.parser, result.error)
    except Exception as exc:
        return ExtractResult("", "parse_failed", "word-com", str(exc))


def extract_pdf_text(path: Path) -> ExtractResult:
    try:
        import pdfplumber

        parts: list[str] = []
        with pdfplumber.open(path) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(text.strip())
        text = compact("\n\n".join(parts))
        return ExtractResult(
            text,
            "parsed" if text else "empty_text",
            "pdfplumber",
            page_count=page_count,
        )
    except Exception as exc:
        return ExtractResult("", "parse_failed", "pdfplumber", str(exc))


def extract_text(path: Path) -> ExtractResult:
    ext = path.suffix.lower()
    if ext in {".md", ".markdown", ".mdx", ".txt"}:
        return read_text_with_fallback(path)
    if ext == ".docx":
        return extract_docx_text(path)
    if ext == ".doc":
        return extract_doc_text(path)
    if ext == ".pdf":
        return extract_pdf_text(path)
    return ExtractResult("", "unsupported", "none", f"unsupported extension: {ext}")


def is_publish_law_notice(name: str) -> bool:
    return "通知" in name and contains_any(
        name,
        (
            "规则",
            "规程",
            "细则",
            "办法",
            "规范",
            "协议",
            "指引",
            "管理",
        ),
    )


def is_form_like_name(name: str) -> bool:
    return contains_any(name, FORM_NAME_TERMS)


def classify(path: Path, extracted: ExtractResult) -> ClassifyResult:
    name = path.name
    name_text = name.lower()
    text = extracted.text or ""
    text_head = text[:20000]

    form_hits = matched_terms(name, FORM_NAME_TERMS)
    general_hits = matched_terms(name, GENERAL_NAME_TERMS)
    manual_strong_hits = matched_terms(name_text, MANUAL_STRONG_NAME_TERMS)
    manual_weak_hits = matched_terms(name_text, MANUAL_WEAK_NAME_TERMS)
    laws_strong_hits = matched_terms(name, LAWS_STRONG_NAME_TERMS)
    laws_name_hits = matched_terms(name, LAWS_NAME_TERMS)

    law_structure_count = count_terms(text_head, LAW_STRUCTURE_TERMS)
    law_normative_count = count_terms(text_head, LAW_NORMATIVE_TERMS)
    manual_action_count = count_terms(text_head, MANUAL_ACTION_TERMS)

    evidence: list[str] = []
    review_reasons: list[str] = []

    if extracted.status == "parse_failed":
        review_reasons.append(f"正文抽取失败：{extracted.parser}")
    elif extracted.status == "empty_text":
        if path.suffix.lower() == ".pdf":
            review_reasons.append("PDF 正文为空，可能为扫描版")
        else:
            review_reasons.append("正文抽取为空")
    elif extracted.status == "unsupported":
        review_reasons.append("不支持的文件类型")

    conflict = False

    # 表单、清单、要素表等优先归 general，除非文件名明确是规则/协议/规范本体。
    if form_hits and not laws_strong_hits:
        evidence.append(f"文件名命中表单/清单类：{'、'.join(form_hits[:4])}")
        confidence = "高" if extracted.status == "parsed" else "中"
        return finalize("general", confidence, evidence, review_reasons)

    if form_hits and laws_strong_hits:
        conflict = True
        review_reasons.append("文件名同时命中表单类和制度类信号")

    if laws_strong_hits:
        evidence.append(f"文件名命中强制度类：{'、'.join(laws_strong_hits[:4])}")
        if law_structure_count or law_normative_count:
            evidence.append(
                f"正文含法规结构/规范表达：结构{law_structure_count}处，规范表达{law_normative_count}处"
            )
        confidence = "高" if extracted.status == "parsed" and not conflict else "中"
        return finalize("laws", confidence, evidence, review_reasons)

    if "指引" in name and not manual_strong_hits and not form_hits:
        evidence.append("文件名命中制度类指引")
        if law_structure_count or law_normative_count:
            evidence.append(
                f"正文含法规结构/规范表达：结构{law_structure_count}处，规范表达{law_normative_count}处"
            )
        confidence = "高" if extracted.status == "parsed" and (law_structure_count or law_normative_count) else "中"
        return finalize("laws", confidence, evidence, review_reasons)

    if is_publish_law_notice(name):
        evidence.append("文件名为发布/修订/调整规则类通知")
        confidence = "高" if extracted.status == "parsed" else "中"
        return finalize("laws", confidence, evidence, review_reasons)

    if laws_name_hits and not general_hits:
        evidence.append(f"文件名命中制度类：{'、'.join(laws_name_hits[:5])}")
        if law_structure_count or law_normative_count:
            evidence.append(
                f"正文含法规结构/规范表达：结构{law_structure_count}处，规范表达{law_normative_count}处"
            )
            confidence = "高" if extracted.status == "parsed" else "中"
        else:
            confidence = "中"
        return finalize("laws", confidence, evidence, review_reasons)

    if law_structure_count >= 2 or law_normative_count >= 8:
        evidence.append(
            f"正文法规结构/规范表达明显：结构{law_structure_count}处，规范表达{law_normative_count}处"
        )
        confidence = "高" if extracted.status == "parsed" else "中"
        return finalize("laws", confidence, evidence, review_reasons)

    if manual_strong_hits:
        evidence.append(f"文件名命中强手册/指南类：{'、'.join(manual_strong_hits[:4])}")
        if manual_action_count:
            evidence.append(f"正文含操作动作词约{manual_action_count}处")
        # 新闻稿式“资料/正式发布”只给中置信度，方便人工确认是否为手册正文。
        confidence = "中" if contains_any(name, NEWS_HINT_TERMS) else "高"
        if extracted.status != "parsed":
            confidence = "中"
        return finalize("manual", confidence, evidence, review_reasons)

    if manual_weak_hits and manual_action_count >= 8:
        evidence.append(f"文件名命中手册/指南弱信号：{'、'.join(manual_weak_hits[:4])}")
        evidence.append(f"正文含操作动作词约{manual_action_count}处")
        confidence = "高" if extracted.status == "parsed" else "中"
        return finalize("manual", confidence, evidence, review_reasons)

    if manual_weak_hits and not contains_any(name, NEWS_HINT_TERMS):
        evidence.append(f"文件名命中手册/指南弱信号：{'、'.join(manual_weak_hits[:4])}")
        confidence = "中"
        return finalize("manual", confidence, evidence, review_reasons)

    if general_hits:
        evidence.append(f"文件名命中普通说明/材料类：{'、'.join(general_hits[:4])}")
        confidence = "高" if extracted.status == "parsed" else "中"
        return finalize("general", confidence, evidence, review_reasons)

    if manual_action_count >= 12 and law_structure_count == 0:
        evidence.append(f"正文以操作动作词为主，约{manual_action_count}处")
        confidence = "中"
        return finalize("manual", confidence, evidence, review_reasons)

    if laws_name_hits:
        evidence.append(f"文件名命中制度类弱信号：{'、'.join(laws_name_hits[:5])}")
        confidence = "中"
        return finalize("laws", confidence, evidence, review_reasons)

    evidence.append("未命中手册或制度类强信号，按普通类型处理")
    confidence = "低" if extracted.status != "parsed" else "中"
    return finalize("general", confidence, evidence, review_reasons)


def finalize(
    category: str,
    confidence: str,
    evidence: list[str],
    review_reasons: list[str],
) -> ClassifyResult:
    if confidence == "低":
        review_reasons.append("低置信度分类")
    return ClassifyResult(
        category=category,
        confidence=confidence,
        evidence="；".join(evidence),
        needs_review="是" if review_reasons else "否",
        review_reason="；".join(dict.fromkeys(review_reasons)),
    )


def classify_files(source_dir: Path) -> list[dict]:
    rows: list[dict] = []
    for path in discover_files(source_dir):
        extracted = extract_text(path)
        result = classify(path, extracted)
        rows.append(
            {
                "文件路径": str(path),
                "文件名": path.name,
                "分类": result.category,
                "置信度": result.confidence,
                "命中依据": result.evidence,
                "是否需要人工复核": result.needs_review,
                "复核原因": result.review_reason,
                "扩展名": path.suffix.lower(),
                "正文抽取状态": extracted.status,
                "正文抽取器": extracted.parser,
                "正文字符数": len(extracted.text or ""),
                "页数": extracted.page_count,
                "抽取错误": extracted.error,
            }
        )
    return rows


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "文件路径",
        "文件名",
        "分类",
        "置信度",
        "命中依据",
        "是否需要人工复核",
        "复核原因",
        "扩展名",
        "正文抽取状态",
        "正文抽取器",
        "正文字符数",
        "页数",
        "抽取错误",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def percent(count: int, total: int) -> str:
    if not total:
        return "0.00%"
    return f"{count / total * 100:.2f}%"


def markdown_table(rows: list[list[str]]) -> list[str]:
    if not rows:
        return []
    header = rows[0]
    out = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
    for row in rows[1:]:
        out.append("| " + " | ".join(str(cell).replace("\n", " ") for cell in row) + " |")
    return out


def write_summary(path: Path, rows: list[dict], source_dir: Path, csv_path: Path) -> None:
    total = len(rows)
    category_counts = Counter(row["分类"] for row in rows)
    confidence_counts = Counter(row["置信度"] for row in rows)
    review_rows = [row for row in rows if row["是否需要人工复核"] == "是"]
    low_rows = [row for row in rows if row["置信度"] == "低"]
    extract_counts = Counter(row["正文抽取状态"] for row in rows)
    ext_counts = Counter(row["扩展名"] for row in rows)

    lines: list[str] = [
        "# 外汇文件文档分类汇总",
        "",
        f"- 源目录：`{source_dir}`",
        f"- 逐文件明细：`{csv_path}`",
        f"- 执行时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- 总文件数：{total}",
        "",
        "## 分类分布",
        "",
    ]
    lines.extend(
        markdown_table(
            [
                ["分类", "数量", "占比"],
                ["manual", str(category_counts.get("manual", 0)), percent(category_counts.get("manual", 0), total)],
                ["laws", str(category_counts.get("laws", 0)), percent(category_counts.get("laws", 0), total)],
                ["general", str(category_counts.get("general", 0)), percent(category_counts.get("general", 0), total)],
            ]
        )
    )
    lines.extend(["", "## 置信度分布", ""])
    lines.extend(
        markdown_table(
            [
                ["置信度", "数量", "占比"],
                ["高", str(confidence_counts.get("高", 0)), percent(confidence_counts.get("高", 0), total)],
                ["中", str(confidence_counts.get("中", 0)), percent(confidence_counts.get("中", 0), total)],
                ["低", str(confidence_counts.get("低", 0)), percent(confidence_counts.get("低", 0), total)],
            ]
        )
    )
    lines.extend(["", "## 正文抽取状态", ""])
    extract_table = [["状态", "数量", "占比"]]
    for status, count in sorted(extract_counts.items()):
        extract_table.append([status, str(count), percent(count, total)])
    lines.extend(markdown_table(extract_table))

    lines.extend(["", "## 文件类型分布", ""])
    ext_table = [["扩展名", "数量", "占比"]]
    for ext, count in sorted(ext_counts.items()):
        ext_table.append([ext, str(count), percent(count, total)])
    lines.extend(markdown_table(ext_table))

    lines.extend(
        [
            "",
            "## 需要人工复核",
            "",
            f"- 需要人工复核文件数：{len(review_rows)}",
            f"- 低置信度文件数：{len(low_rows)}",
            "",
        ]
    )
    review_table = [["文件名", "分类", "置信度", "复核原因"]]
    for row in review_rows[:80]:
        review_table.append([row["文件名"], row["分类"], row["置信度"], row["复核原因"]])
    if len(review_rows) > 80:
        review_table.append([f"... 另有 {len(review_rows) - 80} 个文件", "", "", "详见 CSV"])
    lines.extend(markdown_table(review_table))

    lines.extend(["", "## 低置信度文件", ""])
    if low_rows:
        low_table = [["文件名", "分类", "命中依据"]]
        for row in low_rows:
            low_table.append([row["文件名"], row["分类"], row["命中依据"]])
        lines.extend(markdown_table(low_table))
    else:
        lines.append("无。")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


def write_json_debug(path: Path, rows: list[dict]) -> None:
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Classify foreign exchange documents into manual/laws/general.")
    parser.add_argument("--source", default=str(SOURCE_DIR), help="Source directory.")
    parser.add_argument("--csv", default=str(DEFAULT_CSV), help="Output CSV path.")
    parser.add_argument("--summary", default=str(DEFAULT_SUMMARY), help="Output summary Markdown path.")
    args = parser.parse_args()

    source_dir = Path(args.source)
    csv_path = Path(args.csv)
    summary_path = Path(args.summary)
    if not source_dir.exists():
        raise FileNotFoundError(f"source directory does not exist: {source_dir}")

    rows = classify_files(source_dir)
    write_csv(csv_path, rows)
    write_summary(summary_path, rows, source_dir, csv_path)
    print(
        json.dumps(
            {
                "source_dir": str(source_dir),
                "csv": str(csv_path),
                "summary": str(summary_path),
                "total": len(rows),
                "category_counts": Counter(row["分类"] for row in rows),
                "confidence_counts": Counter(row["置信度"] for row in rows),
                "review_count": sum(1 for row in rows if row["是否需要人工复核"] == "是"),
            },
            ensure_ascii=False,
            default=dict,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
