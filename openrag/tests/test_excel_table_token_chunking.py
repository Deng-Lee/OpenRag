"""Regression tests for token-bounded Excel parsing and chunking."""

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from common.token_utils import num_tokens_from_string
from openrag.chunking.chunk_engine import ChunkEngine
from openrag.chunking.chunk_params import resolve_chunk_method
from openrag.chunking.excel_table_chunker import ExcelChunkingError
from openrag.hierarchy.document_hierarchy_builder import DocumentHierarchyBuilder
from openrag.parsers.adapters.excel_adapter import (
    ExcelParserAdapter,
    StructuredExcelParserAdapter,
)
from openrag.parsers.base import DocumentBlock
from openrag.parsers.factory import ParserFactory


def _excel_block(
    *,
    sheet_name: str = "Data",
    sheet_index: int = 0,
    headers: list[str],
    rows: list[tuple[int, list[str | None]]],
) -> DocumentBlock:
    return DocumentBlock(
        text="preview",
        page=sheet_index + 1,
        offset=0,
        block_type="table",
        block_id=f"xlsx:{sheet_index}:table",
        table_data={
            "schema_version": "excel-table-v1",
            "sheet_name": sheet_name,
            "sheet_index": sheet_index,
            "header_row_index": 1,
            "header": [
                {"column_index": index, "name": name}
                for index, name in enumerate(headers, 1)
            ],
            "rows": [
                {"row_index": row_index, "values": values}
                for row_index, values in rows
            ],
        },
        metadata={
            "source_format": "excel",
            "sheet_name": sheet_name,
            "sheet_index": sheet_index,
        },
    )


def test_excel_chunk_method_isolated_from_other_file_types() -> None:
    assert resolve_chunk_method("book.xlsx", "auto") == "excel_table_token_v1"
    assert resolve_chunk_method("book.XLS", "auto") == "excel_table_token_v1"
    assert resolve_chunk_method("book.xlsx", "txt") == "manual"
    assert resolve_chunk_method("book.csv", "auto") == "manual"
    assert resolve_chunk_method("table.md", "auto") == "manual"
    assert resolve_chunk_method("slides.pptx", "auto") == "presentation"
    assert resolve_chunk_method("report.pdf", "auto") == "pdf_manual"


def test_parser_factory_routes_only_excel_to_structured_adapter() -> None:
    factory = ParserFactory()

    assert isinstance(factory.get_parser("book.xlsx"), StructuredExcelParserAdapter)
    assert isinstance(factory.get_parser("book.xls"), StructuredExcelParserAdapter)
    assert type(factory.get_parser("book.csv")) is ExcelParserAdapter


def test_structured_parser_preserves_excel_coordinates_and_values(tmp_path) -> None:
    path = tmp_path / "structured.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Requirements"
    worksheet.append([None, None, None])
    worksheet.append(["ID", "Enabled", "Due"])
    worksheet.append([0, False, date(2026, 8, 31)])
    worksheet["F50"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    workbook.save(path)

    adapter = StructuredExcelParserAdapter()
    blocks = adapter.parse(str(path))

    assert adapter.parser_version == "excel-structured-v2"
    assert len(blocks) == 1
    block = blocks[0]
    assert block.block_type == "table"
    assert block.metadata["source_format"] == "excel"
    assert block.table_data["schema_version"] == "excel-table-v1"
    assert block.table_data["sheet_name"] == "Requirements"
    assert block.table_data["header"] == [
        {"column_index": 1, "name": "ID"},
        {"column_index": 2, "name": "Enabled"},
        {"column_index": 3, "name": "Due"},
    ]
    assert block.table_data["rows"] == [
        {"row_index": 3, "values": ["0", "False", "2026-08-31T00:00:00"]}
    ]


def test_structured_parser_keeps_header_only_sheet(tmp_path) -> None:
    path = tmp_path / "header-only.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Headers"
    sheet.append(["Account", "Amount"])
    workbook.save(path)

    blocks = StructuredExcelParserAdapter().parse(str(path))
    chunks = ChunkEngine().chunk(
        blocks,
        chunk_size=64,
        chunk_overlap=10,
        chunk_method="excel_table_token_v1",
    )

    assert len(blocks) == 1
    assert blocks[0].table_data["rows"] == []
    assert chunks
    assert all(chunk.metadata["row_start"] == 1 for chunk in chunks)
    assert "Account" in "\n".join(chunk.text for chunk in chunks)


def test_engine_packs_complete_rows_without_exceeding_token_limit() -> None:
    block = _excel_block(
        headers=["ID", "Description"],
        rows=[
            (2, ["REQ-001", "Alpha"]),
            (3, ["REQ-002", "Beta"]),
            (4, ["REQ-003", "Gamma"]),
        ],
    )

    chunks = ChunkEngine().chunk(
        [block],
        chunk_size=80,
        chunk_overlap=10,
        chunk_method="excel_table_token_v1",
    )

    assert len(chunks) == 1
    assert chunks[0].metadata["row_start"] == 2
    assert chunks[0].metadata["row_end"] == 4
    assert all(num_tokens_from_string(chunk.text) <= 80 for chunk in chunks)
    assert all(chunk.metadata["token_count"] <= 80 for chunk in chunks)


def test_engine_splits_wide_row_by_columns() -> None:
    block = _excel_block(
        headers=[f"Column {index}" for index in range(1, 13)],
        rows=[(2, [f"value-{index}-" + "x" * 30 for index in range(1, 13)])],
    )

    chunks = ChunkEngine().chunk(
        [block],
        chunk_size=80,
        chunk_overlap=10,
        chunk_method="excel_table_token_v1",
    )

    assert len(chunks) > 1
    assert all(chunk.metadata["row_start"] == 2 for chunk in chunks)
    assert all(chunk.metadata["column_start"] is not None for chunk in chunks)
    assert all(num_tokens_from_string(chunk.text) <= 80 for chunk in chunks)


def test_engine_splits_long_cell_without_losing_content() -> None:
    body = "商品业务说明" * 400
    block = _excel_block(headers=["Details"], rows=[(7, [body])])

    chunks = ChunkEngine().chunk(
        [block],
        chunk_size=80,
        chunk_overlap=10,
        chunk_method="excel_table_token_v1",
    )

    assert len(chunks) > 1
    assert all(chunk.metadata["fragment_index"] for chunk in chunks)
    assert all(num_tokens_from_string(chunk.text) <= 80 for chunk in chunks)
    recovered = "".join(chunk.text.split("\n\n", 1)[1] for chunk in chunks)
    assert recovered == body


def test_engine_preserves_header_that_exceeds_token_limit() -> None:
    header = "超长列名" * 120
    block = _excel_block(headers=[header], rows=[(2, ["value"])])

    chunks = ChunkEngine().chunk(
        [block],
        chunk_size=80,
        chunk_overlap=10,
        chunk_method="excel_table_token_v1",
    )

    header_chunks = [
        chunk for chunk in chunks
        if chunk.metadata.get("fragment_kind") == "header"
    ]
    recovered_header = "".join(
        chunk.text.split("\n\n", 1)[1] for chunk in header_chunks
    )
    assert recovered_header == header
    assert any(chunk.metadata.get("fragment_kind") == "value" for chunk in chunks)
    assert all(num_tokens_from_string(chunk.text) <= 80 for chunk in chunks)


def test_invalid_excel_table_data_fails_without_retry() -> None:
    block = _excel_block(headers=["ID"], rows=[(2, ["1"])])
    block.table_data["schema_version"] = "unknown"

    with pytest.raises(ExcelChunkingError) as exc_info:
        ChunkEngine().chunk(
            [block],
            chunk_size=80,
            chunk_overlap=10,
            chunk_method="excel_table_token_v1",
        )

    assert exc_info.value.code == "EXCEL_TABLE_DATA_INVALID"
    assert exc_info.value.retryable is False


def test_excel_factory_chunker_and_hierarchy_are_connected(tmp_path) -> None:
    path = tmp_path / "end-to-end.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["ID", "Details"])
    first.append(["1", "说明" * 500])
    second = workbook.create_sheet("Second")
    second.append(["ID", "Status"])
    second.append(["2", "Ready"])
    workbook.save(path)

    parser = ParserFactory().get_parser(str(path))
    blocks = parser.parse(str(path))
    chunks = ChunkEngine().chunk(
        blocks,
        chunk_size=100,
        chunk_overlap=10,
        chunk_method=resolve_chunk_method(str(path), "auto"),
    )
    hierarchy = DocumentHierarchyBuilder().build_hierarchy(chunks)

    assert all(num_tokens_from_string(chunk.text) <= 100 for chunk in chunks)
    assert "First" in hierarchy.l1
    assert "Second" in hierarchy.l1


def test_excel_hierarchy_lists_every_sheet_instead_of_first_50_chunks() -> None:
    chunks = []
    for sheet_index, sheet_name in enumerate(["Summary", "Requirements"]):
        for chunk_index in range(55):
            chunks.append(
                type(
                    "ExcelChunk",
                    (),
                    {
                        "text": f"{sheet_name} row {chunk_index}",
                        "level": 0,
                        "start_offset": chunk_index,
                        "end_offset": chunk_index + 1,
                        "metadata": {
                            "source_format": "excel",
                            "sheet_name": sheet_name,
                            "sheet_index": sheet_index,
                            "row_start": chunk_index + 2,
                            "row_end": chunk_index + 2,
                            "column_start": 1,
                            "column_end": 3,
                        },
                    },
                )()
            )

    hierarchy = DocumentHierarchyBuilder().build_hierarchy(chunks)

    assert "Summary" in hierarchy.l1
    assert "Requirements" in hierarchy.l1
    assert "chunks/0055.md" in hierarchy.l1
