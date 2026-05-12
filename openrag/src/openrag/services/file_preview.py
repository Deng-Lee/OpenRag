"""从 MinIO 中的办公文档字节生成安全、受限的 HTML/纯文本预览（供前端模态窗展示）。"""

from __future__ import annotations

import html
import json
from io import BytesIO
from typing import Literal, Optional, Tuple

PreviewFormat = Literal["html", "text"]

_MAX_XLSX_ROWS = 500
_MAX_XLSX_COLS = 64


def _mime_from_extension(filename: str) -> Optional[str]:
    fn = filename.lower()
    if fn.endswith(".docx"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if fn.endswith(".doc"):
        return "application/msword"
    if fn.endswith(".pptx"):
        return "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    if fn.endswith(".ppt"):
        return "application/vnd.ms-powerpoint"
    if fn.endswith(".xlsx"):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fn.endswith(".xls"):
        return "application/vnd.ms-excel"
    if fn.endswith(".csv"):
        return "text/csv"
    if fn.endswith(".md"):
        return "text/markdown"
    if fn.endswith(".json"):
        return "application/json"
    return None


def build_file_preview(
    data: bytes,
    mime_type: Optional[str],
    filename: str,
) -> Tuple[PreviewFormat, str]:
    """
    根据 MIME / 扩展名将常见办公与文本格式转为 HTML 或纯文本。
    旧版二进制 Office（.doc/.ppt/.xls）不在此解析，返回说明性 HTML。
    """
    effective_mime = (mime_type or "").strip().lower() or None
    if not effective_mime:
        effective_mime = _mime_from_extension(filename)

    if effective_mime in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ):
        return _preview_docx(data)

    if effective_mime in (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ):
        return _preview_pptx(data)

    if effective_mime in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ):
        return _preview_xlsx(data)

    if effective_mime == "application/msword":
        return (
            "html",
            "<p>暂不支持旧版 Word（.doc）内联预览，请使用下载或转换为 .docx。</p>",
        )
    if effective_mime == "application/vnd.ms-powerpoint":
        return (
            "html",
            "<p>暂不支持旧版 PowerPoint（.ppt）内联预览，请使用下载或转换为 .pptx。</p>",
        )
    if effective_mime == "application/vnd.ms-excel":
        return (
            "html",
            "<p>暂不支持旧版 Excel（.xls）内联预览，请使用下载或转换为 .xlsx。</p>",
        )

    if effective_mime in (
        "text/plain",
        "text/csv",
        "application/json",
        "text/markdown",
    ):
        return _preview_plain_text(data, effective_mime)

    raise ValueError(f"No preview converter for mime={effective_mime!r}")


def _preview_docx(data: bytes) -> Tuple[PreviewFormat, str]:
    from docx import Document

    doc = Document(BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(f"<p>{html.escape(t)}</p>")
    for table in doc.tables:
        parts.append("<table class='preview-docx-table'>")
        for row in table.rows:
            parts.append("<tr>")
            for cell in row.cells:
                ctext = html.escape((cell.text or "").replace("\n", " "))
                parts.append(f"<td>{ctext}</td>")
            parts.append("</tr>")
        parts.append("</table>")
    body = "".join(parts) if parts else "<p>（文档无可见段落）</p>"
    return "html", f"<div class='preview-docx'>{body}</div>"


def _preview_pptx(data: bytes) -> Tuple[PreviewFormat, str]:
    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text.strip())
        merged = html.escape("\n".join(t for t in texts if t))
        parts.append(
            f"<section class='preview-ppt-slide'><h3>幻灯片 {i}</h3>"
            f"<div style='white-space:pre-wrap'>{merged or '（无文本）'}</div></section>"
        )
    body = "".join(parts) if parts else "<p>（演示文稿无幻灯片）</p>"
    return "html", f"<div class='preview-pptx'>{body}</div>"


def _preview_xlsx(data: bytes) -> Tuple[PreviewFormat, str]:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_html: list[str] = []
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if row_count >= _MAX_XLSX_ROWS:
                break
            cells = row[:_MAX_XLSX_COLS]
            tds = []
            for c in cells:
                val = "" if c is None else str(c)
                tds.append(f"<td>{html.escape(val)}</td>")
            rows_html.append(f"<tr>{''.join(tds)}</tr>")
            row_count += 1
        truncated_note = ""
        if row_count >= _MAX_XLSX_ROWS:
            truncated_note = (
                "<p class='preview-xlsx-note'>"
                f"仅展示前 {_MAX_XLSX_ROWS} 行，完整内容请下载文件。</p>"
            )
        table = (
            "<table class='preview-xlsx-table' border='1' "
            "style='border-collapse:collapse;width:100%'>"
            f"{''.join(rows_html)}</table>"
        )
        return "html", f"<div class='preview-xlsx'>{table}{truncated_note}</div>"
    finally:
        wb.close()


def _preview_plain_text(data: bytes, mime: str) -> Tuple[PreviewFormat, str]:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("utf-8", errors="replace")

    if mime == "application/json":
        try:
            parsed = json.loads(text)
            text = json.dumps(parsed, ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            pass

    return "text", text
